import pandas as pd
import openpyxl

print("=== Checking Old eDNA Data vs New eDNA Data ===")
old_edna_path = "data/eDNA data_ALL_20260225.xlsx"
new_edna_path = "data/eDNA data_ALL_20260528.xlsx"

xl_old_edna = openpyxl.load_workbook(old_edna_path, read_only=True)
print("Old eDNA sheets:", xl_old_edna.sheetnames)
xl_new_edna = openpyxl.load_workbook(new_edna_path, read_only=True)
print("New eDNA sheets:", xl_new_edna.sheetnames)

df_old_edna = pd.read_excel(old_edna_path, sheet_name=0)
df_new_edna = pd.read_excel(new_edna_path, sheet_name=0)

print(f"Old eDNA shape: {df_old_edna.shape}")
print(f"New eDNA shape: {df_new_edna.shape}")

print("\neDNA Column comparison:")
print("In old but not new:", set(df_old_edna.columns) - set(df_new_edna.columns))
print("In new but not old:", set(df_new_edna.columns) - set(df_old_edna.columns))

print("\n=== Checking Old Cull Data vs New Cull/Manta/RHIS Data ===")
old_cull_path = "data/260201_COTS-Cull-Data-Ewels.xlsx"
new_cull_path = "data/260529-COTS-Manta-Cull-RHIS-Lawrence-CSIRO.xlsx"

xl_old_cull = openpyxl.load_workbook(old_cull_path, read_only=True)
print("Old Cull sheets:", xl_old_cull.sheetnames)
xl_new_cull = openpyxl.load_workbook(new_cull_path, read_only=True)
print("New Cull sheets:", xl_new_cull.sheetnames)

df_old_cull = pd.read_excel(old_cull_path, sheet_name="Cull")
print(f"Old Cull shape: {df_old_cull.shape}")
print("Old Cull columns:", df_old_cull.columns.tolist()[:15])

for sheet in xl_new_cull.sheetnames:
    df_s = pd.read_excel(new_cull_path, sheet_name=sheet, nrows=5)
    print(f"\nNew Cull sheet '{sheet}' shape: nrows sample 5, cols: {len(df_s.columns)}")
    print(f"Sheet '{sheet}' columns:", df_s.columns.tolist()[:15])
